#!/usr/bin/env python3
"""Lista de precios CSV/XLSX → cotización orientativa, sin importes supuestos."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import webbrowser
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from html import escape
from pathlib import Path
from zipfile import BadZipFile
from zoneinfo import ZoneInfo

CENTAVO = Decimal("0.01")


def numero(valor, campo: str, *, positivo: bool = False) -> Decimal:
    try:
        n = Decimal(str(valor).strip())
    except (InvalidOperation, ValueError):
        raise ValueError(f"{campo}: escribe un número decimal válido") from None
    if not n.is_finite() or n < 0 or (positivo and n == 0):
        raise ValueError(
            f"{campo}: el número debe ser finito y {'mayor que cero' if positivo else 'no negativo'}"
        )
    return n


def cargar_precios(ruta: Path) -> list[dict]:
    if ruta.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        try:
            wb = load_workbook(ruta, read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException) as error:
            raise ValueError("El archivo Excel no es válido") from error
        try:
            filas = iter(wb.active.iter_rows(values_only=True))
            cabecera = next(filas, None)
            if not cabecera:
                raise ValueError("La lista de precios está vacía")
            cols = [str(c or "").strip().lower() for c in cabecera]
            precios = [
                dict(zip(cols, r, strict=False)) for r in filas if any(v is not None for v in r)
            ]
        finally:
            wb.close()
    else:
        with ruta.open(encoding="utf-8-sig", newline="") as f:
            lector = csv.DictReader(f)
            cols = [c.strip().lower() for c in (lector.fieldnames or [])]
            precios = [
                {k.strip().lower(): v for k, v in fila.items() if k is not None} for fila in lector
            ]
    if not {"sku", "descripcion"}.issubset(cols) or not {"precio", "precio_unitario"}.intersection(
        cols
    ):
        raise ValueError("La lista requiere sku, descripcion y precio_unitario (o precio)")
    if not precios:
        raise ValueError("La lista de precios está vacía")
    return precios


def buscar(precios: list[dict], clave: str) -> dict | None:
    clave = clave.strip().casefold()
    if not clave:
        raise ValueError("La clave de producto no puede estar vacía")
    exactos = [p for p in precios if clave == str(p.get("sku") or "").strip().casefold()]
    tokens = clave.split()
    candidatos = exactos or [
        p
        for p in precios
        if all(t in f"{p.get('sku') or ''} {p.get('descripcion') or ''}".casefold() for t in tokens)
    ]
    if len(candidatos) > 1:
        raise ValueError(f"Producto ambiguo: {clave}. Usa un SKU único")
    return candidatos[0] if candidatos else None


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--precios", type=Path, required=True)
    ap.add_argument("--pedir", required=True)
    ap.add_argument("--cliente", default="<Cliente>")
    ap.add_argument("--salida", type=Path, default=Path("cotizacion.html"))
    ap.add_argument(
        "--iva", help="Porcentaje confirmado por el cliente; si falta no se calcula el total"
    )
    ap.add_argument("--vigencia", default="<vigencia por confirmar>")
    ap.add_argument("--sin-abrir", action="store_true")
    args = ap.parse_args()
    try:
        precios = cargar_precios(args.precios)
        lineas, subtotal, faltantes = [], Decimal(0), []
        for item in args.pedir.split(","):
            if not item.strip():
                continue
            clave, separador, cantidad = item.partition(":")
            cant = numero(cantidad if separador else "1", "Cantidad", positivo=True)
            p = buscar(precios, clave)
            if not p:
                faltantes.append(clave.strip())
                continue
            valor = p.get("precio_unitario")
            if valor is None or str(valor).strip() == "":
                valor = p.get("precio")
            pu = numero(valor, f"Precio de {clave}")
            importe = (pu * cant).quantize(CENTAVO, rounding=ROUND_HALF_UP)
            subtotal += importe
            lineas.append(
                (p.get("descripcion") or clave, cant, p.get("unidad") or "<unidad>", pu, importe)
            )
        if not lineas and not faltantes:
            raise ValueError("Escribe al menos un producto")
        tasa = numero(args.iva, "IVA") if args.iva is not None else None
        if tasa is not None and tasa > 100:
            raise ValueError("IVA: el porcentaje debe estar entre 0 y 100")
    except (ValueError, OSError, InvalidOperation) as error:
        ap.error(str(error))
    filas = "\n".join(
        f"<tr><td>{escape(str(d))}</td><td>{c:g} {escape(str(u))}</td><td>${pu:,.2f}</td><td>${imp:,.2f}</td></tr>"
        for d, c, u, pu, imp in lineas
    )
    incompleta = bool(faltantes)
    if tasa is not None and not incompleta:
        iva = (subtotal * tasa / 100).quantize(CENTAVO, rounding=ROUND_HALF_UP)
        totales = f"<tr><td colspan=3>IVA ({tasa:g}%)</td><td>${iva:,.2f}</td></tr><tr><td colspan=3>Total</td><td>${subtotal + iva:,.2f}</td></tr>"
    else:
        totales = (
            "<tr><td colspan=3>Total</td><td>&lt;pendiente: confirmar IVA y productos&gt;</td></tr>"
        )
    aviso = (
        "<p>Faltan productos: " + escape(", ".join(faltantes)) + ". Cotización incompleta.</p>"
        if incompleta
        else ""
    )
    html = f"""<!doctype html><html lang="es-MX"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>Cotización {escape(args.cliente)}</title>
<style>body{{font-family:system-ui;max-width:720px;margin:40px auto;padding:0 16px;color:#1a1a1a}}h1{{color:#e63329}}table{{width:100%;border-collapse:collapse}}th,td{{padding:9px;border-bottom:1px solid #ddd;text-align:left}}tfoot{{font-weight:600}}</style>
<h1>Cotización orientativa</h1><p>Cliente: {escape(args.cliente)}<br>Fecha: {dt.datetime.now(ZoneInfo("America/Mexico_City")).date():%d/%m/%Y}<br>Vigencia: {escape(args.vigencia)}</p>{aviso}
<table><thead><tr><th>Concepto</th><th>Cantidad</th><th>P. unitario</th><th>Importe</th></tr></thead><tbody>{filas}</tbody><tfoot><tr><td colspan=3>{"Subtotal de productos encontrados" if incompleta else "Subtotal"}</td><td>${subtotal:,.2f}</td></tr>{totales}</tfoot></table>
<p>Generado por Mecha · Precios orientativos hasta su confirmación.</p></html>"""
    args.salida.write_text(html, encoding="utf-8")
    print(f"→ {args.salida.resolve()}")
    if not args.sin_abrir:
        webbrowser.open(args.salida.resolve().as_uri())
    return 1 if incompleta else 0


if __name__ == "__main__":
    raise SystemExit(main())
